# Build supplementary tables (.xlsx + .docx, three-line format) into out/tables/
import os, glob, math
import pandas as pd
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

R = 'out'; MID = f'{R}/mid'; ARC = f'{R}/causal'; OUT = 'out/tables'
os.makedirs(OUT, exist_ok=True)
TNR = 'Times New Roman'
ORGS = ['kidney', 'liver', 'heart', 'lung', 'colon', 'skin', 'artery', 'esophagus', 'brain']
LBL = {'kidney': 'Kidney', 'liver': 'Liver', 'heart': 'Heart', 'lung': 'Lung', 'colon': 'Colon',
       'skin': 'Skin', 'artery': 'Artery', 'esophagus': 'Esophagus', 'brain': 'Brain',
       'adipose': 'Adipose', 'endometrium': 'Endometrium'}


def three_line(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        x = OxmlElement(f'w:{e}'); x.set(qn('w:val'), 'none'); b.append(x)
    tblPr.append(b)
    def cb(cell, edge):
        tcPr = cell._tc.get_or_add_tcPr(); bs = OxmlElement('w:tcBorders')
        x = OxmlElement(f'w:{edge}'); x.set(qn('w:val'), 'single'); x.set(qn('w:sz'), '8'); x.set(qn('w:color'), '000000')
        bs.append(x); tcPr.append(bs)
    for c in table.rows[0].cells: cb(c, 'top'); cb(c, 'bottom')
    for c in table.rows[-1].cells: cb(c, 'bottom')


def val(x):
    if x is None: return ''
    if isinstance(x, float) and math.isnan(x): return ''
    return str(x)


def write(df, name, caption, sheets=None):
    """sheets: dict{sheetname: df} for a multi-sheet xlsx; docx is written per table in sequence."""
    if sheets:
        with pd.ExcelWriter(f'{OUT}/{name}.xlsx') as w:
            for sn, d in sheets.items(): d.to_excel(w, sheet_name=sn[:31], index=False)
        seq = list(sheets.values())
    else:
        df.to_excel(f'{OUT}/{name}.xlsx', index=False); seq = [df]
    doc = Document(); st = doc.styles['Normal']; st.font.name = TNR; st.font.size = Pt(10)
    p = doc.add_paragraph(); r = p.add_run(caption); r.bold = True; r.font.name = TNR; r.font.size = Pt(10)
    for d in seq:
        t = doc.add_table(rows=1, cols=len(d.columns))
        for j, c in enumerate(d.columns):
            cell = t.rows[0].cells[j]; cell.text = str(c)
            for rr in cell.paragraphs[0].runs: rr.bold = True; rr.font.name = TNR; rr.font.size = Pt(8)
        for _, row in d.iterrows():
            cells = t.add_row().cells
            for j, c in enumerate(d.columns):
                cells[j].text = val(row[c])
                for rr in cells[j].paragraphs[0].runs: rr.font.name = TNR; rr.font.size = Pt(8)
        three_line(t); doc.add_paragraph()
    doc.save(f'{OUT}/{name}.docx'); print('saved', name, [len(x) for x in seq])


S1 = pd.DataFrame([
    ('SenCat', 'secretory (SASP consensus)', sum(1 for l in open('raw/sencat_symbols.csv', encoding='utf-8') if l.strip()) - 1, 'raw/sencat_symbols.csv'),
    ('SenMayo', 'secretory (SASP)', sum(1 for l in open('raw/sig_SenMayo.txt') if l.strip()), 'raw/sig_SenMayo.txt'),
    ('Fridman', 'senescence-up', sum(1 for l in open('raw/sig_fridman.txt') if l.strip()), 'raw/sig_fridman.txt'),
    ('Reactome-SASP', 'secretory (SASP)', sum(1 for l in open('raw/sig_reactome_sasp.txt') if l.strip()), 'raw/sig_reactome_sasp.txt'),
    ('Reactome cellular senescence', 'cell-cycle', sum(1 for l in open('raw/sig_reactome.txt') if l.strip()), 'raw/sig_reactome.txt'),
    ('GO:0090398', 'cell-cycle', sum(1 for l in open('raw/sig_go_senescence.txt') if l.strip()), 'raw/sig_go_senescence.txt'),
], columns=['Signature', 'Type', 'Genes', 'Source file'])
write(S1, 'TableS3', 'Six transcriptional senescence signatures and their sizes.')

rows = []
for f in sorted(glob.glob(f'{MID}/organ_up_*.txt')):
    o = os.path.basename(f)[9:-4]
    if o not in ORGS: continue
    for l in open(f):
        g = l.strip()
        if g: rows.append({'Organ': LBL.get(o, o), 'Gene': g})
S2 = pd.DataFrame(rows)
write(S2, 'TableS1', 'Disease signature genes (consensus up-regulated) per organ.')

ps = pd.read_csv(f'{R}/pathway_shared_ge3organs.csv'); ps.insert(0, 'Gene set', 'shared (>=3 organs)')
po = pd.read_csv(f'{R}/pathway_organ_specific.csv'); po.insert(0, 'Gene set', 'organ-specific')
pa = pd.read_csv(f'{R}/pathway_all_interface.csv'); pa.insert(0, 'Gene set', 'all interface')
S3 = pd.concat([ps, po, pa], ignore_index=True)
write(S3, 'TableS4', 'Complete pathway enrichment results.')

S4 = pd.read_csv(f'{R}/subtype_consistency_per_gene.csv')
write(S4, 'TableS2', 'Cross-subtype consistency per interface gene.')

COH = {'kidney': 'GSE30529', 'liver': 'GSE25097', 'heart': 'GSE79962', 'lung': 'GSE24206', 'colon': 'GSE16879',
       'skin': 'GSE14905', 'artery': 'GSE28829', 'esophagus': 'GSE26886', 'brain': 'GSE48350', 'brain28146': 'GSE28146'}
S5 = pd.read_csv(f'{R}/valid_enrichment.csv')
S5.insert(2, 'Cohort', S5['organ'].map(COH))
S5['organ'] = S5['organ'].map(lambda o: LBL.get(o, o))
write(S5, 'TableS5', 'Independent validation cohorts and enrichment of the senescence consensus.')

S6 = {
    'MR_eQTL':         pd.read_csv(f'{ARC}/mr_eqtl_results.csv'),
    'MR_pQTL':         pd.read_csv(f'{ARC}/mr_pqtl_results.csv'),
    'MR_screen':       pd.read_csv(f'{ARC}/screen_mr_results.csv'),
    'Coloc_full':      pd.read_csv(f'{ARC}/coloc_results.csv'),
    'Coloc_narrow':    pd.read_csv(f'{ARC}/coloc_narrow_results.csv'),
    'MR_tissue_eQTL':  pd.read_csv(f'{ARC}/eqtl_tissue_mr.csv'),
}
with pd.ExcelWriter(f'{OUT}/TableS6.xlsx') as w:
    for sn, d in S6.items(): d.to_excel(w, sheet_name=sn[:31], index=False)
doc = Document(); st = doc.styles['Normal']; st.font.name = TNR; st.font.size = Pt(10)
p = doc.add_paragraph(); r = p.add_run('Causal analysis of the shared interface genes.')
r.bold = True; r.font.name = TNR; r.font.size = Pt(10)
for sn, d in S6.items():
    ph = doc.add_paragraph(); rr = ph.add_run(sn); rr.italic = True; rr.font.name = TNR; rr.font.size = Pt(9)
    t = doc.add_table(rows=1, cols=len(d.columns))
    for j, c in enumerate(d.columns):
        cell = t.rows[0].cells[j]; cell.text = str(c)
        for q in cell.paragraphs[0].runs: q.bold = True; q.font.name = TNR; q.font.size = Pt(8)
    for _, row in d.iterrows():
        cells = t.add_row().cells
        for j, c in enumerate(d.columns):
            cells[j].text = val(row[c])
            for q in cells[j].paragraphs[0].runs: q.font.name = TNR; q.font.size = Pt(8)
    three_line(t); doc.add_paragraph()
doc.save(f'{OUT}/TableS6.docx'); print('saved TableS6', {k: len(v) for k, v in S6.items()})
print('DONE')
